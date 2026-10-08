"""Spreadsheet parser service for bulk certificate generation.

Parses Excel (.xlsx) and CSV (.csv) files into validated recipient lists.
Includes intelligent column header detection, date normalization, and
Pydantic validation.
"""

import csv
import io
from datetime import datetime, date
from typing import Any

from openpyxl import load_workbook
from pydantic import ValidationError

from app.schemas import RecipientCreate


# Header aliases for forgiving column detection
NAME_ALIASES = {"name", "recipient_name", "recipient", "full_name", "participant_name", "student_name", "student", "candidate"}
COURSE_ALIASES = {"course_name", "course", "workshop", "event", "program", "training", "event_name", "title", "course_title"}
DATE_ALIASES = {"date", "completion_date", "issue_date", "event_date", "certified_date", "certificate_date"}
EMAIL_ALIASES = {"email", "recipient_email", "mail", "e-mail", "email_address"}


def _normalize_key(key: str) -> str:
    """Normalize a column header for matching."""
    return str(key).strip().lower().replace(" ", "_").replace("-", "_")


def _find_column_mapping(headers: list[str]) -> dict[str, str]:
    """Map detected header names to standardized field names.
    
    Returns a dict mapping standard field ('name', 'course_name', 'date', 'email')
    to the actual column header in the sheet.
    """
    mapping = {}
    normalized_headers = {h: _normalize_key(h) for h in headers if h is not None}

    for original, norm in normalized_headers.items():
        if norm in NAME_ALIASES and "name" not in mapping:
            mapping["name"] = original
        elif norm in COURSE_ALIASES and "course_name" not in mapping:
            mapping["course_name"] = original
        elif norm in DATE_ALIASES and "date" not in mapping:
            mapping["date"] = original
        elif norm in EMAIL_ALIASES and "email" not in mapping:
            mapping["email"] = original

    missing = []
    if "name" not in mapping:
        missing.append("Name")
    if "course_name" not in mapping:
        missing.append("Course / Event Name")
    if "date" not in mapping:
        missing.append("Date")

    if missing:
        raise ValueError(
            f"Could not find required column(s): {', '.join(missing)}. "
            f"Found columns: {list(normalized_headers.keys())}. "
            "Please ensure your spreadsheet contains columns for Name, Course, and Date."
        )

    return mapping


def _format_date(val: Any) -> str:
    """Convert spreadsheet date value to YYYY-MM-DD string format."""
    if isinstance(val, (datetime, date)):
        return val.strftime("%Y-%m-%d")
    
    val_str = str(val).strip()
    if not val_str:
        raise ValueError("Date cannot be empty")
    
    # Try standard YYYY-MM-DD
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y", "%m-%d-%Y", "%Y/%m/%d", "%d.%m.%Y", "%B %d, %Y"):
        try:
            parsed = datetime.strptime(val_str, fmt)
            return parsed.strftime("%Y-%m-%d")
        except ValueError:
            pass

    return val_str


def parse_csv_data(content: bytes) -> list[RecipientCreate]:
    """Parse CSV content into a list of validated RecipientCreate models."""
    # Decode with fallback for encoding
    text = ""
    for encoding in ("utf-8-sig", "utf-8", "latin-1", "cp1252"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue

    if not text:
        raise ValueError("Failed to decode CSV file. Please ensure it is saved in UTF-8 format.")

    reader = csv.reader(io.StringIO(text))
    rows = [row for row in reader if any(cell.strip() for cell in row)]
    if not rows:
        raise ValueError("CSV file is empty.")

    headers = rows[0]
    mapping = _find_column_mapping(headers)

    header_indices = {headers[i]: i for i in range(len(headers))}
    name_idx = header_indices[mapping["name"]]
    course_idx = header_indices[mapping["course_name"]]
    date_idx = header_indices[mapping["date"]]
    email_idx = header_indices.get(mapping.get("email"))

    recipients: list[RecipientCreate] = []
    errors: list[str] = []

    for row_num, row in enumerate(rows[1:], start=2):
        if not any(cell.strip() for cell in row):
            continue

        try:
            name_val = row[name_idx].strip() if name_idx < len(row) else ""
            course_val = row[course_idx].strip() if course_idx < len(row) else ""
            raw_date = row[date_idx].strip() if date_idx < len(row) else ""
            date_val = _format_date(raw_date)

            email_val = None
            if email_idx is not None and email_idx < len(row):
                raw_email = row[email_idx].strip()
                email_val = raw_email if raw_email else None

            recipients.append(
                RecipientCreate(
                    name=name_val,
                    course_name=course_val,
                    date=date_val,
                    email=email_val,
                )
            )
        except (ValueError, ValidationError) as e:
            errors.append(f"Row {row_num}: {e}")

    if errors and len(recipients) == 0:
        raise ValueError(f"Failed to parse CSV recipients: {errors[:5]}")

    return recipients


def parse_excel_data(content: bytes) -> list[RecipientCreate]:
    """Parse Excel (.xlsx) file into a list of validated RecipientCreate models."""
    try:
        wb = load_workbook(io.BytesIO(content), data_only=True)
    except Exception as e:
        raise ValueError(f"Could not read Excel file: {e}")

    ws = wb.active
    if not ws:
        raise ValueError("Excel file contains no active sheet.")

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError("Excel worksheet is empty.")

    # Find the header row (first row with non-empty cells)
    header_row_idx = -1
    for idx, row in enumerate(rows):
        if any(cell is not None and str(cell).strip() != "" for cell in row):
            header_row_idx = idx
            break

    if header_row_idx == -1:
        raise ValueError("Excel worksheet contains no data rows.")

    headers = [str(cell) if cell is not None else "" for cell in rows[header_row_idx]]
    mapping = _find_column_mapping(headers)

    header_indices = {headers[i]: i for i in range(len(headers))}
    name_idx = header_indices[mapping["name"]]
    course_idx = header_indices[mapping["course_name"]]
    date_idx = header_indices[mapping["date"]]
    email_idx = header_indices.get(mapping.get("email"))

    recipients: list[RecipientCreate] = []
    errors: list[str] = []

    for row_num, row in enumerate(rows[header_row_idx + 1:], start=header_row_idx + 2):
        if not any(cell is not None and str(cell).strip() != "" for cell in row):
            continue

        try:
            name_val = str(row[name_idx]).strip() if name_idx < len(row) and row[name_idx] is not None else ""
            course_val = str(row[course_idx]).strip() if course_idx < len(row) and row[course_idx] is not None else ""
            raw_date = row[date_idx] if date_idx < len(row) else ""
            date_val = _format_date(raw_date)

            email_val = None
            if email_idx is not None and email_idx < len(row) and row[email_idx] is not None:
                raw_email = str(row[email_idx]).strip()
                email_val = raw_email if raw_email else None

            recipients.append(
                RecipientCreate(
                    name=name_val,
                    course_name=course_val,
                    date=date_val,
                    email=email_val,
                )
            )
        except (ValueError, ValidationError) as e:
            errors.append(f"Row {row_num}: {e}")

    if errors and len(recipients) == 0:
        raise ValueError(f"Failed to parse Excel recipients: {errors[:5]}")

    return recipients


def parse_recipients_file(content: bytes, filename: str) -> list[RecipientCreate]:
    """Parse recipients from either an Excel (.xlsx) or CSV (.csv) file."""
    ext = filename.lower().split(".")[-1] if "." in filename else ""
    if ext == "csv":
        return parse_csv_data(content)
    elif ext in ("xlsx", "xlsm"):
        return parse_excel_data(content)
    else:
        raise ValueError(f"Unsupported spreadsheet format '.{ext}'. Supported formats are: .xlsx, .csv")
